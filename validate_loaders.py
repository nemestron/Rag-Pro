import os
import logging
from typing import List
import docx
from fpdf import FPDF

# Configure basic logging to see the output
logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')

from ingestion.loaders import load_all_sources

# --- Mock Streamlit UploadedFile ---
class MockUploadedFile:
    def __init__(self, name, content_bytes):
        self.name = name
        self._content = content_bytes
        
    def getvalue(self):
        return self._content

# --- Create Sample Files ---
def create_sample_files():
    samples_dir = os.path.join("data", "samples")
    os.makedirs(samples_dir, exist_ok=True)
    
    # 1. Create Sample TXT
    txt_path = os.path.join(samples_dir, "sample_document.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("This is a sample text document for validating the RAG Pro TXT loader.")
        
    # 2. Create Sample DOCX
    docx_path = os.path.join(samples_dir, "sample_document.docx")
    doc = docx.Document()
    doc.add_paragraph("This is a sample Word document for validating the RAG Pro DOCX loader.")
    doc.save(docx_path)
    
    # 3. Create Sample PDF
    pdf_path = os.path.join(samples_dir, "sample_document.pdf")
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="This is a sample PDF document for validating the RAG Pro PDF loader.", ln=1, align='L')
    pdf.output(pdf_path)
    
    return txt_path, docx_path, pdf_path

# --- Run Validation ---
def run_validation():
    print("\n--- Generating Sample Files ---")
    txt_path, docx_path, pdf_path = create_sample_files()
    
    print("\n--- Mocking Streamlit Uploads ---")
    with open(txt_path, "rb") as f:
        mock_txt = MockUploadedFile("sample_document.txt", f.read())
        
    with open(docx_path, "rb") as f:
        mock_docx = MockUploadedFile("sample_document.docx", f.read())
        
    with open(pdf_path, "rb") as f:
        mock_pdf = MockUploadedFile("sample_document.pdf", f.read())

    print("\n--- Executing load_all_sources ---")
    test_urls = ["https://en.wikipedia.org/wiki/Retrieval-augmented_generation"]
    test_raw_text = "This is a direct raw text input for validation."
    
    try:
        documents = load_all_sources(
            urls=test_urls,
            pdf_files=[mock_pdf],
            docx_files=[mock_docx],
            txt_files=[mock_txt],
            raw_text=test_raw_text
        )
        
        print(f"\n✅ SUCCESS: Extracted {len(documents)} total document chunks.")
        print("\n--- Content Snippets & Metadata ---")
        for i, doc in enumerate(documents):
            print(f"\nDocument {i+1}:")
            print(f"Metadata: {doc.metadata}")
            print(f"Content Snippet: {doc.page_content[:150]}...")
            
    except Exception as e:
        print(f"\n❌ FAILED validation: {e}")

if __name__ == "__main__":
    # Note: Requires 'fpdf' to generate the dummy PDF. Run: pip install fpdf
    run_validation()
